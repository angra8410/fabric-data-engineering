import os
import csv
import openpyxl

raw_dir = 'data/raw'
files = sorted(os.listdir(raw_dir))

report = []

for f in files:
    path = os.path.join(raw_dir, f)
    size_kb = os.path.getsize(path) / 1024
    header = f"\n{'='*75}\nFILE: {f} ({size_kb:.1f} KB)\n{'='*75}"
    report.append(header)
    
    if f.endswith('.csv'):
        try:
            with open(path, 'r', encoding='utf-8', errors='ignore') as fp:
                reader = csv.reader(fp)
                rows = [next(reader, None) for _ in range(6)]
                total_lines = sum(1 for _ in fp) + len([r for r in rows if r is not None])
            header_row = rows[0] if rows and rows[0] else []
            report.append(f"Type: CSV | Approximate rows: {total_lines}")
            report.append(f"Columns ({len(header_row)}): {header_row[:15]}")
            if len(header_row) > 15:
                report.append(f"  ... and {len(header_row) - 15} more columns")
            if len(rows) > 1 and rows[1]:
                report.append(f"Sample Row 1: {rows[1][:15]}")
        except Exception as e:
            report.append(f"Error reading CSV: {e}")
            
    elif f.endswith('.xlsx'):
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            report.append(f"Type: Excel (.xlsx) | Sheet count: {len(wb.sheetnames)}")
            report.append(f"Sheets: {wb.sheetnames}")
            for sname in wb.sheetnames:
                ws = wb[sname]
                sample_rows = []
                for r in ws.iter_rows(max_row=5, values_only=True):
                    sample_rows.append(r)
                report.append(f"  --- Sheet: '{sname}' ---")
                for r_idx, r in enumerate(sample_rows, 1):
                    non_empty = [str(x)[:30] if x is not None else '' for x in r[:12]]
                    if any(non_empty):
                        report.append(f"    R{r_idx}: {non_empty}")
            wb.close()
        except Exception as e:
            report.append(f"Error reading Excel: {e}")

output_text = "\n".join(report)
print(output_text[:4000])
with open("audit_summary.txt", "w", encoding="utf-8") as out:
    out.write(output_text)
print("\n[Full audit report saved to audit_summary.txt]")

import csv
import openpyxl

print("=== 1. OFM Population Final (Skagit) ===")
with open('data/raw/ofm_april1_population_final(Population).csv', 'r', encoding='utf-8', errors='ignore') as f:
    r = csv.reader(f)
    for row in r:
        if len(row) > 3 and row[2].strip() == 'Skagit':
            print("  ", row[3].strip(), "-->", row[4:])

print("\n=== 2. OFM Permits 1990-present (Skagit) ===")
wb = openpyxl.load_workbook('data/raw/ofm_april1_postcensal_permits_1990-present.xlsx', read_only=True, data_only=True)
ws = wb['PermitCompletionDemolition']
cities = set()
years = set()
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] and str(r[0]).strip().lower() == 'skagit':
        cities.add(str(r[1]).strip())
        years.add(str(r[2]).strip())
print(f"  Cities in Skagit ({len(cities)}):", sorted(list(cities)))
print(f"  Years covered ({len(years)}): min {min(years)}, max {max(years)}")
wb.close()

print("\n=== 3. SAEP UGA Population (Skagit) ===")
with open('data/raw/saep_uga20p(Total Population).csv', 'r', encoding='utf-8', errors='ignore') as f:
    r = csv.reader(f)
    uga_names = []
    for row in r:
        if len(row) > 2 and row[0].strip().lower() == 'skagit':
            uga_names.append(row[2].strip())
print(f"  UGAs in Skagit ({len(uga_names)}):", sorted(uga_names))

print("\n=== 4. UGA 2025 Permits (Skagit County) ===")
with open('data/raw/UGA_Dev_Permits_2025.csv', 'r', encoding='utf-8', errors='ignore') as f:
    r = csv.reader(f)
    next(r) # header
    permit_ugas = set()
    work_classes = set()
    permit_types = set()
    for row in r:
        if len(row) > 5:
            permit_ugas.add(row[5].strip())
            work_classes.add(row[2].strip())
            permit_types.add(row[1].strip())
print("  Permit UGAs:", sorted(list(permit_ugas)))
print("  Permit Work Classes:", sorted(list(work_classes)))
print("  Permit Types:", sorted(list(permit_types)))

print("\n=== 5. 2045 Projections & Allocations ===")
wb = openpyxl.load_workbook('data/raw/2045-Adopted-SkagitCounty-GrowthProjectionsAndAllocationsTables.xlsx', data_only=True)
ws = wb['Table 1']
for r in ws.iter_rows(min_row=5, max_row=20, values_only=True):
    name = str(r[0]).strip() if r[0] is not None else ''
    if name:
        p2022 = r[5]
        p2045 = r[7]
        growth = r[10]
        print(f"  {name:25} | 2022: {p2022} | 2045 Alloc: {p2045} | Projected Growth: {growth}")
wb.close()

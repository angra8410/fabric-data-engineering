import os
import csv
import openpyxl

def audit_file(filepath, name):
    print(f"\n==================== {name} ====================")
    if filepath.endswith('.csv'):
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            reader = csv.reader(f)
            lines = [next(reader, None) for _ in range(15)]
        for i, l in enumerate(lines):
            if l:
                print(f"L{i+1}: {l[:8]}")
    elif filepath.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
        for s in wb.sheetnames[:3]:
            print(f"--- Sheet: {s} ---")
            ws = wb[s]
            for i, r in enumerate(ws.iter_rows(max_row=8, values_only=True)):
                non_empty = [str(x)[:20] if x is not None else '' for x in r[:8]]
                print(f"R{i+1}: {non_empty}")
        wb.close()

files = [
    ("2045 Allocations", "data/raw/2045-Adopted-SkagitCounty-GrowthProjectionsAndAllocationsTables.xlsx"),
    ("UGA 2025 Permits", "data/raw/UGA_Dev_Permits_2025.csv"),
    ("OFM Pop Final", "data/raw/ofm_april1_population_final(Population).csv"),
    ("QCEW 2025", "data/raw/2025-QCEW-annual-averages-revised(Skagit County) (1).csv"),
    ("OFM Permits 1990-pres", "data/raw/ofm_april1_postcensal_permits_1990-present.xlsx"),
    ("SAEP UGA Pop", "data/raw/saep_uga20p(Total Population).csv")
]

for name, path in files:
    audit_file(path, name)

import os
import re
import csv
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

RAW_DIR = "data/raw"
PROCESSED_DIR = "data/processed"
os.makedirs(PROCESSED_DIR, exist_ok=True)

# -----------------------------------------------------------------------------
# JURISDICTION NORMALIZATION MAP
# -----------------------------------------------------------------------------
JURISDICTION_MASTER = [
    {"Jurisdiction_ID": "JUR-01", "Jurisdiction_Name": "Anacortes", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.5126, "Longitude": -122.6127},
    {"Jurisdiction_ID": "JUR-02", "Jurisdiction_Name": "Burlington", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.4757, "Longitude": -122.3254},
    {"Jurisdiction_ID": "JUR-03", "Jurisdiction_Name": "Concrete", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.5393, "Longitude": -121.7471},
    {"Jurisdiction_ID": "JUR-04", "Jurisdiction_Name": "Hamilton", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.5246, "Longitude": -121.9868},
    {"Jurisdiction_ID": "JUR-05", "Jurisdiction_Name": "La Conner", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.3943, "Longitude": -122.4957},
    {"Jurisdiction_ID": "JUR-06", "Jurisdiction_Name": "Lyman", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.5251, "Longitude": -122.0621},
    {"Jurisdiction_ID": "JUR-07", "Jurisdiction_Name": "Mount Vernon", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.4212, "Longitude": -122.3340},
    {"Jurisdiction_ID": "JUR-08", "Jurisdiction_Name": "Sedro-Woolley", "Jurisdiction_Type": "Incorporated City/Town", "Is_Incorporated": 1, "Is_UGA": 1, "Latitude": 48.5039, "Longitude": -122.2363},
    {"Jurisdiction_ID": "JUR-09", "Jurisdiction_Name": "Bay View Ridge UGA", "Jurisdiction_Type": "Urban Growth Area (UGA)", "Is_Incorporated": 0, "Is_UGA": 1, "Latitude": 48.4717, "Longitude": -122.4239},
    {"Jurisdiction_ID": "JUR-10", "Jurisdiction_Name": "Swinomish UGA", "Jurisdiction_Type": "Urban Growth Area (UGA)", "Is_Incorporated": 0, "Is_UGA": 1, "Latitude": 48.4069, "Longitude": -122.5186},
    {"Jurisdiction_ID": "JUR-11", "Jurisdiction_Name": "Unincorporated Rural (outside UGAs)", "Jurisdiction_Type": "Unincorporated Rural Area", "Is_Incorporated": 0, "Is_UGA": 0, "Latitude": 48.4800, "Longitude": -121.8000},
]

NAME_NORM_MAP = {
    "anacortes": "Anacortes",
    "burlington": "Burlington",
    "concrete": "Concrete",
    "hamilton": "Hamilton",
    "la conner": "La Conner",
    "laconner": "La Conner",
    "lyman": "Lyman",
    "mount vernon": "Mount Vernon",
    "mount\nvernon": "Mount Vernon",
    "mt vernon": "Mount Vernon",
    "sedro-woolley": "Sedro-Woolley",
    "sedro woolley": "Sedro-Woolley",
    "sedrowoolley": "Sedro-Woolley",
    "sedro-\nwoolley": "Sedro-Woolley",
    "sedro- woolley": "Sedro-Woolley",
    "bay view ridge": "Bay View Ridge UGA",
    "bayview ridge": "Bay View Ridge UGA",
    "bayview\nridge": "Bay View Ridge UGA",
    "bay view ridge uga": "Bay View Ridge UGA",
    "swinomish": "Swinomish UGA",
    "swinomish non-trust lands": "Swinomish UGA",
    "swinomish uga": "Swinomish UGA",
    "unincorporated": "Unincorporated Rural (outside UGAs)",
    "unincorporated skagit county": "Unincorporated Rural (outside UGAs)",
    "unincorporated rural (outside ugas)": "Unincorporated Rural (outside UGAs)",
    "unincorporated rural (outside of ugas)": "Unincorporated Rural (outside UGAs)",
    "rural (outside ugas)": "Unincorporated Rural (outside UGAs)",
    "rural (outside of ugas)": "Unincorporated Rural (outside UGAs)",
    "rural (outside of uga's)": "Unincorporated Rural (outside UGAs)",
    "rural (outside of\nuga's)": "Unincorporated Rural (outside UGAs)",
    "rural (outside of\nugas)": "Unincorporated Rural (outside UGAs)",
    "rural (outside of uga’s)": "Unincorporated Rural (outside UGAs)",
    "rural (outside of\nuga’s)": "Unincorporated Rural (outside UGAs)",
}

def normalize_name(raw_name):
    if not raw_name:
        return None
    cleaned = str(raw_name).strip().lower().replace("  ", " ").replace("’", "'")
    if cleaned in NAME_NORM_MAP:
        return NAME_NORM_MAP[cleaned]
    cleaned_flat = cleaned.replace("\n", " ").replace("- ", "-").replace("  ", " ").strip()
    if cleaned_flat in NAME_NORM_MAP:
        return NAME_NORM_MAP[cleaned_flat]
    if cleaned_flat.startswith("rural") or cleaned_flat.startswith("unincorporated"):
        return "Unincorporated Rural (outside UGAs)"
    return NAME_NORM_MAP.get(cleaned_flat, str(raw_name).strip())

def clean_num(val):
    if val is None or val == "" or str(val).strip() in ["*", "-", "N/A", "None"]:
        return 0
    s = str(val).replace(",", "").replace("$", "").replace("%", "").strip()
    try:
        return float(s) if "." in s else int(s)
    except:
        return 0

# CAI total-employment method: average of the template's yearly (covered + self-employed) / covered
# ratios, 1999-2019 (21 observations). County-wide, all-industry figure: apply to the TOTAL row only.
SELF_EMP_MULTIPLIER = 1.15458

# -----------------------------------------------------------------------------
# 1. BUILD DIM_JURISDICTION
# -----------------------------------------------------------------------------
def build_dim_jurisdiction():
    print("-> Building Dim_Jurisdiction...")
    df = pd.DataFrame(JURISDICTION_MASTER)
    df["County_Name"] = "Skagit"
    df["State"] = "WA"
    return df

# -----------------------------------------------------------------------------
# 2. BUILD DIM_CALENDARYEAR
# -----------------------------------------------------------------------------
def build_dim_calendaryear():
    print("-> Building Dim_CalendarYear...")
    years = list(range(1990, 2046))
    records = []
    for y in years:
        decade = f"{(y // 10) * 10}s"
        if y < 2000:
            cycle = "Early GMA Foundation (1990-1999)"
        elif y < 2020:
            cycle = "Prior Comprehensive Cycle (2000-2019)"
        else:
            cycle = "Active GMA 20-Year Horizon (2020-2045)"
            
        records.append({
            "Year": y,
            "Decade": decade,
            "GMA_Planning_Cycle": cycle,
            "Is_Historical": 1 if y <= 2026 else 0,
            "Is_Prototype_Year_2025": 1 if y == 2025 else 0,
            "Is_Production_Year_2026": 1 if y == 2026 else 0,
            "Is_Target_Year_2045": 1 if y == 2045 else 0
        })
    return pd.DataFrame(records)

# -----------------------------------------------------------------------------
# 3. BUILD DIM_GMA_2045_TARGET
# -----------------------------------------------------------------------------
def build_dim_gma_2045_target(dim_jur):
    print("-> Building Dim_GMA_2045_Target...")
    path = os.path.join(RAW_DIR, "2045-Adopted-SkagitCounty-GrowthProjectionsAndAllocationsTables.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Table 1']
    
    # Table 1: Population Allocations (rows 5 to 14, 16)
    pop_dict = {}
    for r in [5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 16]:
        name = normalize_name(ws.cell(r, 1).value)
        p22 = clean_num(ws.cell(r, 6).value)
        p45 = clean_num(ws.cell(r, 8).value)
        pgrowth = clean_num(ws.cell(r, 11).value)
        pshare = clean_num(ws.cell(r, 13).value)
        pop_dict[name] = (p22, p45, pgrowth, pshare)

    # Table 2: Housing Allocations - Net New Housing Needed 2020-2045 (rows 22 to 31, 33)
    hsg_dict = {}
    for r in [22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 33]:
        name = normalize_name(ws.cell(r, 1).value)
        tot_hsg = clean_num(ws.cell(r, 16).value)
        hsg_dict[name] = tot_hsg

    tot_hsg_table2 = sum(hsg_dict.values())
    assert tot_hsg_table2 == 17450, f"Table 2 total housing must equal 17,450, got {tot_hsg_table2}"
    assert hsg_dict.get("Unincorporated Rural (outside UGAs)") == 3490, f"Rural housing target must equal 3,490, got {hsg_dict.get('Unincorporated Rural (outside UGAs)')}"

    # Table 3: Employment Allocations (rows 40 to 49, 51)
    emp_dict = {}
    for r in [40, 41, 42, 43, 44, 45, 46, 47, 48, 49, 51]:
        name = normalize_name(ws.cell(r, 1).value)
        e22 = clean_num(ws.cell(r, 6).value)
        e45 = clean_num(ws.cell(r, 9).value)
        egrowth = clean_num(ws.cell(r, 11).value)
        eshare = clean_num(ws.cell(r, 13).value)
        emp_dict[name] = (e22, e45, egrowth, eshare)

    wb.close()
    
    records = []
    for _, j in dim_jur.iterrows():
        jname = j["Jurisdiction_Name"]
        p22, p45, pgrowth, pshare = pop_dict.get(jname, (0, 0, 0, 0))
        h45 = hsg_dict.get(jname, 0)
        e22, e45, egrowth, eshare = emp_dict.get(jname, (0, 0, 0, 0))
        records.append({
            "Jurisdiction_ID": j["Jurisdiction_ID"],
            "Jurisdiction_Name": jname,
            "Jurisdiction_Type": j["Jurisdiction_Type"],
            "Baseline_2022_Population": p22,
            "Target_2045_Population": p45,
            "Projected_2045_Population_Growth": pgrowth,
            "Population_Growth_Share_Pct": round(float(pshare), 4) if pshare else 0.0,
            "Target_2045_Housing_Units": h45,
            "Baseline_2022_Employment": e22,
            "Target_2045_Employment": e45,
            "Projected_2045_Employment_Growth": egrowth,
            "Employment_Growth_Share_Pct": round(float(eshare), 4) if eshare else 0.0,
            "CAI_Self_Employment_Multiplier": SELF_EMP_MULTIPLIER,
            "Data_Source": "Appendix A. Skagit County 2045 Growth Projections & Allocations (O20250002)"
        })
        
    return pd.DataFrame(records)

# -----------------------------------------------------------------------------
# 4. BUILD FACT_POPULATION
# -----------------------------------------------------------------------------
def build_fact_population(dim_jur):
    print("-> Building Fact_Population...")
    
    # 4.1 SAEP UGA Population (UGAs historical 2010 to 2026)
    # Read first so we can subtract UGAs from Unincorporated for each year
    saep_path = os.path.join(RAW_DIR, "saep_uga20p(Total Population).csv")
    uga_records = []
    uga_by_year = {}
    with open(saep_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for _ in range(11):
            next(reader, None)
        header = next(reader)
        uga_year_cols = []
        for idx, col in enumerate(header):
            col_clean = col.strip()
            if "change" in col_clean.lower():
                continue
            m = re.search(r'\b(20\d\d)\b', col_clean)
            if m:
                uga_year_cols.append((idx, int(m.group(1))))
                
        for row in reader:
            if len(row) > 2 and row[0].strip().lower() == "skagit":
                raw_uga = row[2].strip()
                norm_uga = normalize_name(raw_uga)
                # Only take unincorporated UGAs not covered by OFM city records (Bay View Ridge & Swinomish)
                if norm_uga in ["Bay View Ridge UGA", "Swinomish UGA"]:
                    for col_idx, yr in uga_year_cols:
                        if col_idx < len(row):
                            pop_val = clean_num(row[col_idx])
                            if pop_val > 0:
                                uga_records.append({
                                    "Jurisdiction_Name": norm_uga,
                                    "Year": yr,
                                    "Population_Count": pop_val,
                                    "Data_Source_Type": "OFM_SAEP_UGA_Estimate"
                                })
                                uga_by_year[yr] = uga_by_year.get(yr, 0) + pop_val

    # 4.2 OFM April 1 Population Final (Cities & Unincorporated County)
    ofm_path = os.path.join(RAW_DIR, "ofm_april1_population_final(Population).csv")
    city_records = []
    uninc_by_year = {}
    with open(ofm_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for _ in range(4): # skip first 4 lines
            next(reader, None)
        header = next(reader)
        year_cols = []
        for idx, col in enumerate(header):
            m = re.search(r'(20\d\d)', col)
            if m:
                year_cols.append((idx, int(m.group(1))))
                
        for row in reader:
            if len(row) > 3 and row[2].strip() == "Skagit":
                raw_jur = row[3].strip()
                if raw_jur in ["Skagit County", "Incorporated Skagit County"]:
                    continue # County rollups happen via measure aggregation in Power BI
                if raw_jur in ["Unincorporated", "Unincorporated Skagit County"]:
                    for col_idx, yr in year_cols:
                        if col_idx < len(row):
                            pop_val = clean_num(row[col_idx])
                            if pop_val > 0:
                                uninc_by_year[yr] = pop_val
                else:
                    norm_jur = normalize_name(raw_jur)
                    for col_idx, yr in year_cols:
                        if col_idx < len(row):
                            pop_val = clean_num(row[col_idx])
                            if pop_val > 0:
                                city_records.append({
                                    "Jurisdiction_Name": norm_jur,
                                    "Year": yr,
                                    "Population_Count": pop_val,
                                    "Data_Source_Type": "OFM_April1_Official_Determination"
                                })

    # 4.3 Derive Unincorporated Rural (outside UGAs) = OFM Unincorporated - UGAs for each year
    rural_records = []
    for yr, uninc_pop in sorted(uninc_by_year.items()):
        uga_pop = uga_by_year.get(yr, 0)
        rural_pop = uninc_pop - uga_pop
        rural_records.append({
            "Jurisdiction_Name": "Unincorporated Rural (outside UGAs)",
            "Year": yr,
            "Population_Count": rural_pop,
            "Data_Source_Type": "Derived (OFM Unincorporated minus UGAs)"
        })

    # Combine all granular non-overlapping entities
    records = city_records + uga_records + rural_records
    df = pd.DataFrame(records)
    # Deduplicate if any overlap
    df = df.drop_duplicates(subset=["Jurisdiction_Name", "Year"])
    
    # Merge Jurisdiction_ID
    df = pd.merge(df, dim_jur[["Jurisdiction_ID", "Jurisdiction_Name"]], on="Jurisdiction_Name", how="inner")
    
    # Calculate Prior Year & YoY Changes
    df = df.sort_values(by=["Jurisdiction_ID", "Year"]).reset_index(drop=True)
    df["Prior_Year_Population"] = df.groupby("Jurisdiction_ID")["Population_Count"].shift(1)
    df["YoY_Population_Change"] = df["Population_Count"] - df["Prior_Year_Population"]
    df["YoY_Growth_Rate_Pct"] = (df["YoY_Population_Change"] / df["Prior_Year_Population"]).round(4)
    
    df["Fact_Population_Key"] = [f"FPOP-{i+1:05d}" for i in range(len(df))]
    
    # Verification check: For 2025, verify 81,220 (8 cities) + 4,278 (2 UGAs) + 49,102 (rural) == 134,600
    df_2025 = df[df["Year"] == 2025]
    cities_2025 = int(df_2025[df_2025["Data_Source_Type"] == "OFM_April1_Official_Determination"]["Population_Count"].sum())
    ugas_2025 = int(df_2025[df_2025["Data_Source_Type"] == "OFM_SAEP_UGA_Estimate"]["Population_Count"].sum())
    rural_2025 = int(df_2025[df_2025["Jurisdiction_Name"] == "Unincorporated Rural (outside UGAs)"]["Population_Count"].sum())
    total_2025 = int(df_2025["Population_Count"].sum())
    print(f"   [ETL Population Verification 2025] Cities: {cities_2025:,} | UGAs: {ugas_2025:,} | Rural: {rural_2025:,} | Total: {total_2025:,}")
    assert cities_2025 == 81220, f"Expected cities 81,220, got {cities_2025}"
    assert ugas_2025 == 4278, f"Expected UGAs 4,278, got {ugas_2025}"
    assert rural_2025 == 49102, f"Expected rural 49,102, got {rural_2025}"
    assert total_2025 == 134600, f"Expected total 134,600, got {total_2025}"

    cols_order = [
        "Fact_Population_Key", "Jurisdiction_ID", "Year", "Jurisdiction_Name",
        "Population_Count", "Prior_Year_Population", "YoY_Population_Change", "YoY_Growth_Rate_Pct",
        "Data_Source_Type"
    ]
    return df[cols_order]

# -----------------------------------------------------------------------------
# 5. BUILD FACT_HOUSINGPERMITS
# -----------------------------------------------------------------------------
def build_fact_housing_permits(dim_jur):
    print("-> Building Fact_HousingPermits...")
    records = []
    
    # 5.1 Historical Permits from OFM (1990 - 2026)
    ofm_perm_path = os.path.join(RAW_DIR, "ofm_april1_postcensal_permits_1990-present.xlsx")
    wb = openpyxl.load_workbook(ofm_perm_path, data_only=True)
    ws = wb['PermitCompletionDemolition']
    
    for r in ws.iter_rows(min_row=2, values_only=True):
        county = str(r[0]).strip().lower() if r[0] else ""
        if county == "skagit":
            raw_city = str(r[1]).strip()
            norm_jur = normalize_name(raw_city)
            year = int(r[2])
            
            sf_perm = clean_num(r[3])
            sf_comp = clean_num(r[4])
            sf_demo = clean_num(r[5])
            
            dup_perm = clean_num(r[6])
            dup_comp = clean_num(r[7])
            dup_demo = clean_num(r[8])
            
            mf34_perm = clean_num(r[9])
            mf34_comp = clean_num(r[10])
            mf34_demo = clean_num(r[11])
            
            mf5_perm = clean_num(r[12])
            mf5_comp = clean_num(r[13])
            mf5_demo = clean_num(r[14])
            
            adu_perm = clean_num(r[15])
            adu_comp = clean_num(r[16])
            adu_demo = clean_num(r[17])
            
            mob_perm = 0
            mob_comp = 0
            mob_demo = 0
            
            tot_perm = sf_perm + dup_perm + mf34_perm + mf5_perm + adu_perm
            tot_comp = sf_comp + dup_comp + mf34_comp + mf5_comp + adu_comp
            tot_demo = sf_demo + dup_demo + mf34_demo + mf5_demo + adu_demo
            net_new = tot_perm - tot_demo
            
            records.append({
                "Jurisdiction_Name": norm_jur,
                "Year": year,
                "Single_Family_Units": sf_perm,
                "Duplex_Units": dup_perm,
                "MultiFamily_3_4_Units": mf34_perm,
                "MultiFamily_5_Plus_Units": mf5_perm,
                "ADU_Units": adu_perm,
                "Mobile_Home_Units": mob_perm,
                "Total_Permitted_Units": tot_perm,
                "Completed_Units": tot_comp,
                "Demolished_Units": tot_demo,
                "Net_New_Units": net_new,
                "Total_Valuation_USD": 0.0,
                "Data_Source": "OFM Postcensal Housing Permits (1990-present)"
            })
    wb.close()
    
    # 5.2 2025 UGA Permits File (Add Bay View Ridge and Swinomish specifically for 2025)
    uga_perm_path = os.path.join(RAW_DIR, "UGA_Dev_Permits_2025.csv")
    with open(uga_perm_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)
        uga_agg = {}
        for row in reader:
            raw_uga = row.get("Select the Urban Growth Area of the parcel", "").strip()
            if not raw_uga:
                continue
            norm_uga = normalize_name(raw_uga)
            wclass = row.get("Permit Work Class", "").strip().lower()
            val = clean_num(row.get("Permit Valuation", 0))
            
            if norm_uga not in uga_agg:
                uga_agg[norm_uga] = {"sf": 0, "adu": 0, "add": 0, "tot": 0, "val": 0.0}
            
            if "single family" in wclass:
                uga_agg[norm_uga]["sf"] += 1
                uga_agg[norm_uga]["tot"] += 1
            elif "accessory dwelling" in wclass or "adu" in wclass:
                uga_agg[norm_uga]["adu"] += 1
                uga_agg[norm_uga]["tot"] += 1
            else:
                uga_agg[norm_uga]["add"] += 1
            uga_agg[norm_uga]["val"] += val
            
        for u_name, data in uga_agg.items():
            if u_name in ["Bay View Ridge UGA", "Swinomish UGA"]:
                records.append({
                    "Jurisdiction_Name": u_name,
                    "Year": 2025,
                    "Single_Family_Units": data["sf"],
                    "Duplex_Units": 0,
                    "MultiFamily_3_4_Units": 0,
                    "MultiFamily_5_Plus_Units": 0,
                    "ADU_Units": data["adu"],
                    "Mobile_Home_Units": 0,
                    "Total_Permitted_Units": data["tot"],
                    "Completed_Units": data["tot"],
                    "Demolished_Units": 0,
                    "Net_New_Units": data["tot"],
                    "Total_Valuation_USD": round(float(data["val"]), 2),
                    "Data_Source": "Skagit County UGA Building Permits Log (2025)"
                })

    df = pd.DataFrame(records)
    df = pd.merge(df, dim_jur[["Jurisdiction_ID", "Jurisdiction_Name"]], on="Jurisdiction_Name", how="inner")
    df = df.sort_values(by=["Jurisdiction_ID", "Year"]).reset_index(drop=True)
    df["Fact_Housing_Key"] = [f"FHSG-{i+1:05d}" for i in range(len(df))]
    
    cols_order = [
        "Fact_Housing_Key", "Jurisdiction_ID", "Year", "Jurisdiction_Name",
        "Single_Family_Units", "Duplex_Units", "MultiFamily_3_4_Units", "MultiFamily_5_Plus_Units",
        "ADU_Units", "Mobile_Home_Units", "Total_Permitted_Units", "Completed_Units",
        "Demolished_Units", "Net_New_Units", "Total_Valuation_USD", "Data_Source"
    ]
    return df[cols_order]

# -----------------------------------------------------------------------------
# 6. BUILD DIM_CAI_EMPLOYMENT_BENCHMARK
# -----------------------------------------------------------------------------
def build_dim_cai_employment_benchmark():
    """Historical CAI series (observed years) + the derived 2022 baseline row.

    Observed rows (1999-2019) are read from the CAI template. The 2022 row is NOT an
    observation: Census NES data in the template stops at 2019, so 2022 total employment is
    covered employment x the average ratio. NES and ratio are left blank for that row.
    """
    print("-> Building Dim_CAI_Employment_Benchmark...")
    cai_path = os.path.join(RAW_DIR, "CAI.Total Employment Calc Template DRAFT.2024 0206.xlsx")
    wb = openpyxl.load_workbook(cai_path, data_only=True)
    ws = wb['Total Employment Calculation']

    # Fail loudly if the template's average ratio no longer matches the constant used in the fact table.
    template_avg = float(ws["D36"].value)
    if round(template_avg, 5) != SELF_EMP_MULTIPLIER:
        raise ValueError(
            f"CAI template average ratio is {template_avg:.5f} but SELF_EMP_MULTIPLIER is "
            f"{SELF_EMP_MULTIPLIER}. Update the constant (and ADR-012) before running the ETL."
        )

    observed = []
    # Template columns D..X (4..24) hold the observed years (1999-2019).
    for c in range(4, 25):
        yr = ws.cell(20, c).value
        if not isinstance(yr, (int, float)):
            continue
        self_emp = ws.cell(21, c).value
        total_emp = ws.cell(26, c).value
        ratio = ws.cell(31, c).value
        if not all(isinstance(v, (int, float)) for v in (self_emp, total_emp, ratio)):
            continue
        observed.append({
            "Year": int(yr),
            "Covered_Employment_QCEW": int(round(total_emp - self_emp)),
            "Self_Employment_NES": int(round(self_emp)),
            "Total_Employment_Combined": int(round(total_emp)),
            "Self_Employment_Ratio": round(float(ratio), 5),
            "Is_Observed": 1,
            "Data_Source": "CAI Total Employment Calc Template / BLS QCEW & Census NES (observed)",
        })

    ratios = [r["Self_Employment_Ratio"] for r in observed]
    all_years_avg = round(sum(ratios) / len(ratios), 5)
    last10 = ratios[-10:]
    last10_avg = round(sum(last10) / len(last10), 5)

    # Derived baseline year (template cells D15 = year, D16 = covered, D41 = total estimate).
    baseline_year = int(ws["D15"].value)
    baseline = {
        "Year": baseline_year,
        "Covered_Employment_QCEW": int(ws["D16"].value),
        "Self_Employment_NES": None,          # not observed: NES stops at 2019 in the template
        "Total_Employment_Combined": int(round(ws["D41"].value)),
        "Self_Employment_Ratio": None,        # derived with the average ratio, not observed
        "Is_Observed": 0,
        "Data_Source": (f"Derived: {baseline_year} covered employment x average ratio "
                        f"(CAI template D41). Adopted Appendix A baseline is 59,571."),
    }
    wb.close()

    df = pd.DataFrame(observed + [baseline]).sort_values("Year").reset_index(drop=True)
    df["Ratio_Average_All_Years"] = all_years_avg     # 21 observations, 1999-2019 (= the multiplier)
    df["Ratio_Average_Last_10_Obs"] = last10_avg      # 2010-2019, shown for comparison only
    df["Benchmark_Key"] = [f"CAI-EMP-{y}" for y in df["Year"]]
    for col in ("Covered_Employment_QCEW", "Self_Employment_NES", "Total_Employment_Combined", "Is_Observed"):
        df[col] = df[col].astype("Int64")
    cols_order = [
        "Benchmark_Key", "Year", "Covered_Employment_QCEW", "Self_Employment_NES",
        "Total_Employment_Combined", "Self_Employment_Ratio", "Ratio_Average_All_Years",
        "Ratio_Average_Last_10_Obs", "Is_Observed", "Data_Source",
    ]
    return df[cols_order]

# -----------------------------------------------------------------------------
# 7. BUILD FACT_EMPLOYMENT
# -----------------------------------------------------------------------------
_ROLLUP_TITLES = {"total": "TOTAL", "government": "GOV", "not elsewhere classified": "NEC"}
_GOV_SUBROWS = {"federal government": "GOV-FED", "state government": "GOV-STATE", "local government": "GOV-LOCAL"}


def parse_qcew_value(val):
    """Parse an ESD QCEW cell. Blank or suppressed ('*') cells return None, never 0,
    so suppression is not confused with a true zero."""
    if val is None:
        return None
    s = str(val).strip()
    if s in ("", "*", "-", "N/A", "None"):
        return None
    s = s.replace(",", "").replace("$", "").strip()
    try:
        return float(s) if "." in s else int(s)
    except ValueError:
        return None


def read_qcew_rows(path):
    """Read an ESD QCEW file into normalized rows with a stable Industry_Key.

    The header row is located by content (not by a fixed line count), so the TOTAL row
    directly below it is always kept. Rollup rows are normalized so both years use the same
    codes: Total -> 'TOTAL', Government -> 'GOV', Not Elsewhere Classified -> 'NEC'.
    """
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        rows = list(csv.reader(f))
    hdr = next(i for i, r in enumerate(rows) if r and r[0].strip().lower().startswith("2-digit naics"))

    out, parent_2d = [], None
    for r in rows[hdr + 1:]:
        if len(r) < 4:
            continue
        n2, n3, title = r[0].strip(), r[1].strip(), r[2].strip()
        if not title:
            continue                                   # blank spacer and footnote rows
        tkey = title.lower()
        if not n2 and not n3 and tkey in _ROLLUP_TITLES:
            n2 = _ROLLUP_TITLES[tkey]                  # 2026 file leaves these codes blank

        if n2:
            parent_2d = n2
            level, key = ("County Total", "TOTAL") if n2 == "TOTAL" else ("2-digit", n2)
        elif n3:
            level, key = "3-digit", n3
        elif tkey in _GOV_SUBROWS:
            level, key = "Government sub-sector", _GOV_SUBROWS[tkey]
        else:
            level, key = "Residual", f"{parent_2d}-OTHER"   # the 'Other industries' lines
        out.append({"n2": n2, "n3": n3, "title": title, "key": key, "level": level,
                    "parent_2d": parent_2d if key != "TOTAL" else "TOTAL", "cells": r})
    return out


def build_fact_employment():
    """Fact_Employment: ESD QCEW covered employment by industry.

    - Industry_Key is stable across years, so year-over-year joins do not depend on titles.
    - Suppressed ('*') cells are blank (null) and flagged, not 0.
    - Estimated_Total_Employment (covered x CAI ratio) exists ONLY on the county TOTAL row. The
      ratio is a county-wide, all-industry figure and is not valid for individual sectors.
    - 2026 is Q1-only: Annual_Average_Employment is blank, months Apr-Dec are blank, and
      Q1_Average_Employment is populated for both years so Q1 can be compared with Q1.
    """
    print("-> Building Fact_Employment...")
    annual_path = os.path.join(RAW_DIR, "2025-QCEW-annual-averages-revised(Skagit County) (1).csv")
    q1_path = os.path.join(RAW_DIR, "2026Q1-QCEW-preliminary(Skagit County).csv")
    month_cols = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

    annual_rows = read_qcew_rows(annual_path)
    canonical_title = {r["key"]: r["title"] for r in annual_rows}   # 2025 titles are the reference

    def base_record(r, year, status, period):
        return {
            "Year": year, "Period_Type": period, "Industry_Key": r["key"], "Row_Level": r["level"],
            "Parent_2Digit_Code": r["parent_2d"], "NAICS_2Digit_Code": r["n2"], "NAICS_3Digit_Code": r["n3"],
            "Industry_Subsector_Title": canonical_title.get(r["key"], r["title"]),
            "Title_As_Reported": r["title"], "Data_Status": status,
            "Is_County_Total": 1 if r["key"] == "TOTAL" else 0,
        }

    records = []
    for r in annual_rows:                                            # 2025 annual
        c = r["cells"]
        months = [parse_qcew_value(c[i]) if len(c) > i else None for i in range(4, 16)]
        avg = parse_qcew_value(c[17]) if len(c) > 17 else None
        q1 = round(sum(months[:3]) / 3) if all(m is not None for m in months[:3]) else None
        rec = base_record(r, 2025, "2025 Official Revised Averages", "Annual")
        rec.update({
            "Average_Establishments": parse_qcew_value(c[3]),
            "Annual_Average_Employment": avg,
            "Q1_Average_Employment": q1,
            "Is_Suppressed": 1 if avg is None else 0,
            "Average_Annual_Wage_USD": parse_qcew_value(c[18]) if len(c) > 18 else None,
        })
        rec.update({f"{m}_Employment": v for m, v in zip(month_cols, months)})
        records.append(rec)

    if os.path.exists(q1_path):                                      # 2026 Q1 preliminary
        for r in read_qcew_rows(q1_path):
            c = r["cells"]
            months = [parse_qcew_value(c[i]) if len(c) > i else None for i in range(4, 7)]
            q1 = parse_qcew_value(c[8]) if len(c) > 8 else None
            qwage = parse_qcew_value(c[9]) if len(c) > 9 else None
            rec = base_record(r, 2026, "2026 Q1 Preliminary", "Q1")
            rec.update({
                "Average_Establishments": parse_qcew_value(c[3]),
                "Annual_Average_Employment": None,                   # no annual figure exists yet
                "Q1_Average_Employment": q1,
                "Is_Suppressed": 1 if q1 is None else 0,
                "Average_Annual_Wage_USD": int(qwage * 4) if qwage is not None else None,  # Q1 wage x 4 (approximation)
            })
            rec.update({f"{m}_Employment": (months[i] if i < 3 else None) for i, m in enumerate(month_cols)})
            records.append(rec)

    df = pd.DataFrame(records)

    # CAI total-employment estimate: county total row only.
    is_total = df["Is_County_Total"] == 1
    df["Estimated_Total_Employment"] = (df["Annual_Average_Employment"] * SELF_EMP_MULTIPLIER).round()
    df["Estimated_Total_Employment_Q1"] = (df["Q1_Average_Employment"] * SELF_EMP_MULTIPLIER).round()
    df.loc[~is_total, ["Estimated_Total_Employment", "Estimated_Total_Employment_Q1"]] = None
    df["CAI_Self_Employment_Multiplier"] = None
    df.loc[is_total, "CAI_Self_Employment_Multiplier"] = SELF_EMP_MULTIPLIER

    dup = df.duplicated(subset=["Year", "Industry_Key"], keep=False)
    if dup.any():
        raise ValueError(f"Industry_Key is not unique within a year: {df.loc[dup, ['Year','Industry_Key']].values.tolist()}")

    int_cols = (["Average_Establishments", "Annual_Average_Employment", "Q1_Average_Employment",
                 "Estimated_Total_Employment", "Estimated_Total_Employment_Q1", "Average_Annual_Wage_USD",
                 "Is_Suppressed"] + [f"{m}_Employment" for m in month_cols])
    for col in int_cols:
        df[col] = df[col].round().astype("Int64")

    df["Fact_Employment_Key"] = [f"FEMP-{i+1:04d}" for i in range(len(df))]
    cols_order = [
        "Fact_Employment_Key", "Year", "Period_Type", "Industry_Key", "Row_Level", "Parent_2Digit_Code",
        "NAICS_2Digit_Code", "NAICS_3Digit_Code", "Industry_Subsector_Title", "Title_As_Reported",
        "Is_County_Total", "Is_Suppressed", "Average_Establishments",
        "Annual_Average_Employment", "Q1_Average_Employment",
        "Estimated_Total_Employment", "Estimated_Total_Employment_Q1", "CAI_Self_Employment_Multiplier",
        "Average_Annual_Wage_USD",
    ] + [f"{m}_Employment" for m in month_cols] + ["Data_Status"]
    return df[cols_order]

# -----------------------------------------------------------------------------
# 8. BUILD FACT_HOUSING_AMI
# -----------------------------------------------------------------------------
def build_fact_housing_ami(dim_jur):
    print("-> Building Fact_Housing_AMI...")
    path = os.path.join("data/templates", "Template_Housing_AMI_Master.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Housing_AMI_Master']
    
    records = []
    for r in ws.iter_rows(min_row=9, max_row=28, values_only=True):
        if not r or not r[1]:
            continue
        rec_id = r[1]
        yr = r[2]
        jur = r[3]
        stype = r[4]
        ami_0_30 = clean_num(r[5])
        ami_31_50 = clean_num(r[6])
        ami_51_80 = clean_num(r[7])
        ami_81_100 = clean_num(r[8])
        ami_101_120 = clean_num(r[9])
        ami_120_plus = clean_num(r[10])
        tot_units = clean_num(r[11])
        status = r[12]
        
        records.append({
            "Fact_AMI_Key": rec_id,
            "Year": yr,
            "Jurisdiction_Name": jur,
            "Structure_Type": stype,
            "AMI_0_to_30_Pct_Units": ami_0_30,
            "AMI_31_to_50_Pct_Units": ami_31_50,
            "AMI_51_to_80_Pct_Units": ami_51_80,
            "AMI_81_to_100_Pct_Units": ami_81_100,
            "AMI_101_to_120_Pct_Units": ami_101_120,
            "AMI_Greater_120_Pct_Units": ami_120_plus,
            "Total_AMI_Units": tot_units,
            "Reconciliation_Status": status
        })
    wb.close()
    
    df = pd.DataFrame(records)
    df = pd.merge(df, dim_jur[["Jurisdiction_ID", "Jurisdiction_Name"]], on="Jurisdiction_Name", how="inner")
    cols_order = [
        "Fact_AMI_Key", "Jurisdiction_ID", "Year", "Jurisdiction_Name", "Structure_Type",
        "AMI_0_to_30_Pct_Units", "AMI_31_to_50_Pct_Units", "AMI_51_to_80_Pct_Units",
        "AMI_81_to_100_Pct_Units", "AMI_101_to_120_Pct_Units", "AMI_Greater_120_Pct_Units",
        "Total_AMI_Units", "Reconciliation_Status"
    ]
    return df[cols_order]

# -----------------------------------------------------------------------------
# MAIN ETL EXECUTION
# -----------------------------------------------------------------------------
def run_etl():
    print("="*70)
    print("EXECUTING ETL PIPELINE: SCOG ANNUAL GROWTH MONITORING STAR SCHEMA")
    print("="*70)
    
    # 1. Dimensions
    dim_jur = build_dim_jurisdiction()
    dim_cal = build_dim_calendaryear()
    dim_gma = build_dim_gma_2045_target(dim_jur)
    dim_cai = build_dim_cai_employment_benchmark()
    
    # 2. Facts
    fact_pop = build_fact_population(dim_jur)
    fact_hsg = build_fact_housing_permits(dim_jur)
    fact_emp = build_fact_employment()
    fact_ami = build_fact_housing_ami(dim_jur)
    
    # Save CSVs
    tables = {
        "Dim_Jurisdiction": dim_jur,
        "Dim_CalendarYear": dim_cal,
        "Dim_GMA_2045_Target": dim_gma,
        "Dim_CAI_Employment_Benchmark": dim_cai,
        "Fact_Population": fact_pop,
        "Fact_HousingPermits": fact_hsg,
        "Fact_Employment": fact_emp,
        "Fact_Housing_AMI": fact_ami
    }
    
    print("\n-> Saving Processed CSV Tables:")
    for name, df in tables.items():
        csv_path = os.path.join(PROCESSED_DIR, f"{name}.csv")
        df.to_csv(csv_path, index=False, encoding="utf-8")
        print(f"   [CSV] {name:30} | Rows: {len(df):6d} | Cols: {len(df.columns):2d} -> {csv_path}")

    # Save to Multi-tab Master Excel Model for Power BI
    excel_path = os.path.join(PROCESSED_DIR, "SCOG_Star_Schema_Data_Model.xlsx")
    print(f"\n-> Building Master Excel Workbook for Power BI Desktop: {excel_path}...")
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        for name, df in tables.items():
            df.to_excel(writer, sheet_name=name, index=False)
            
    # Format openpyxl workbook with tables
    wb = openpyxl.load_workbook(excel_path)
    for name in tables.keys():
        ws = wb[name]
        ws.views.sheetView[0].showGridLines = True
        # auto-fit width
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or '')) for c in col)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    wb.save(excel_path)
    print("   [XLSX] Complete Multi-Tab Model Saved Successfully!")

    print("\n" + "="*70)
    print("ETL PIPELINE COMPLETE: STAR SCHEMA READY FOR POWER BI INGESTION")
    print("="*70)

if __name__ == "__main__":
    run_etl()
